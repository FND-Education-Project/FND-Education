"""Browser-based crossword workshop for GitHub Codespaces.

Start with:
    python -m streamlit run tools/puzzles/crossword_app.py \
        --server.address 0.0.0.0 --server.port 8501

The Streamlit process runs **inside** Codespaces; it does not publish a new
service on the FND website. Keep the forwarded port PRIVATE. Nothing is
committed or published automatically.
"""
from __future__ import annotations

from contextlib import redirect_stderr, redirect_stdout
from io import BytesIO, StringIO
import hashlib
import json
import os
import signal
import threading
from pathlib import Path
import shutil
import sys
import tempfile

import pandas as pd
import pymupdf as fitz
import streamlit as st
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from generate_crosswords import main as generate_cli  # noqa: E402
from crossword_builder.deliver import prepare_downloads
from crossword_builder.publish import title_measure, validate_title, restore_layout

HEADERS = ("clue", "answer", "priority", "lesson", "module")
MAX_UPLOAD = 15_000_000


def _read_rows(raw: bytes) -> pd.DataFrame:
    """Read exactly the five source fields; do not silently lose other data."""
    book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        headers = {str(value).strip().lower(): i for i, value in enumerate(first)
                   if value is not None}
        if not set(HEADERS).issubset(headers):
            raise ValueError("Excel needs all five headers: " + ", ".join(HEADERS))
        result = []
        for cells in sheet.iter_rows(min_row=2, values_only=True):
            values = [
                "" if headers[field] >= len(cells) or cells[headers[field]] is None
                else str(cells[headers[field]]).strip()
                for field in HEADERS
            ]
            if any(values):
                result.append(dict(zip(HEADERS, values)))
        return pd.DataFrame(result, columns=list(HEADERS), dtype=str)
    finally:
        book.close()


def _save_rows(original: bytes, edited: pd.DataFrame) -> bytes:
    """Preserve the existing workbook, sheets and formatting while editing rows.

    Only data in the five named columns on the first sheet is changed.
    Literal answer/clue text starting with '=' stays text (not an Excel formula).
    """
    book = load_workbook(BytesIO(original))
    try:
        sheet = book.worksheets[0]
        headers = {
            str(cell.value).strip().lower(): cell.column
            for cell in sheet[1] if cell.value is not None
        }
        missing = set(HEADERS) - set(headers)
        if missing:
            raise ValueError("Missing Excel headers: " + ", ".join(sorted(missing)))
        rows: list[list[str]] = []
        for _, record in edited.iterrows():
            values = [
                "" if pd.isna(record[name]) else str(record[name]).strip()
                for name in HEADERS
            ]
            if any(values):
                rows.append(values)
        if len(rows) > 10_000:
            raise ValueError("The workbook supports at most 10,000 questions.")
        for row in range(2, max(sheet.max_row, len(rows) + 1) + 1):
            for name in HEADERS:
                cell = sheet.cell(row, headers[name])
                cell.value = None
        for row_num, values in enumerate(rows, start=2):
            for field, value in zip(HEADERS, values):
                cell = sheet.cell(row_num, headers[field])
                if field in ("module", "lesson") and value.isdecimal():
                    cell.value = int(value)
                elif field == "priority" and value in ("1", "2", "3"):
                    cell.value = int(value)
                else:
                    cell.value = value
                    # Prevent user-authored text from becoming spreadsheet code.
                    if value.startswith("="):
                        cell.data_type = "s"
        memory = BytesIO()
        book.save(memory)
        return memory.getvalue()
    finally:
        book.close()


@st.cache_data(show_spinner="Preparing your titled PDFs…")
def _finished_pdf_bundle(
    candidate_json: str, selection_json: str, title: str
) -> tuple[str, dict[str, tuple[str, bytes]], bytes]:
    """Export only the selected layout; title changes do not rerun placement."""
    with tempfile.TemporaryDirectory(prefix="fnd-final-puzzle-") as directory:
        return prepare_downloads(
            json.loads(candidate_json), json.loads(selection_json),
            title, Path(directory),
        )


def _thumbnail(path: Path, zoom: float = 1.05) -> bytes:
    """Render a PDF page for a clickable preview without a cloud service."""
    doc = fitz.open(str(path))
    try:
        return doc[0].get_pixmap(matrix=fitz.Matrix(zoom, zoom),
                                alpha=False).tobytes("png")
    finally:
        doc.close()


st.set_page_config(
    page_title="FND Crossword Workshop",
    page_icon="🧩", layout="wide",
)
heading, exit_column = st.columns([7, 2], vertical_alignment="center")
with heading:
    st.title("FND Crossword Workshop")
with exit_column:
    with st.popover("Exit workshop"):
        st.write(
            "Stop the private workshop server and close this session? "
            "Remember to save or download any workbook edits first. "
            "You can restart the workshop from Codespaces later."
        )
        if st.button("Stop server and exit", type="primary"):
            previous = st.session_state.get("crossword_results")
            if previous:
                shutil.rmtree(previous["parent"], ignore_errors=True)
            st.session_state.clear()
            # Finish this response before shutting down Streamlit itself.
            # This stops the private workshop only, not the Codespace.
            threading.Timer(
                1.5, lambda: os.kill(os.getpid(), signal.SIGTERM)
            ).start()
            st.success("Workshop closed. You can close this browser tab.")
            st.stop()
st.caption(
    "Edit your question bank, select a lesson or module, compare five puzzles, "
    "and download the chosen Normal, Easy, and Answer Key PDFs."
)

# 1. The master workbook remains a regular .xlsx file in tools/puzzles/.
st.subheader("1. Question bank")
available = sorted(HERE.glob("*.xlsx"))
source_options = {path.name: path for path in available}
preferred = "FND_Crossword_Question_Bank_Template.xlsx"
names = list(source_options)
default = names.index(preferred) if preferred in names else 0
selected_name = st.selectbox(
    "Use an Excel workbook already in the repository",
    names if names else ["(No Excel file found in tools/puzzles)"],
    index=default,
)
uploaded = st.file_uploader(
    "Or import an Excel workbook from your computer", type=["xlsx"],
    help="Uploading here does not automatically change the GitHub repository.",
)
if uploaded is not None:
    raw = uploaded.getvalue()
    filename = Path(uploaded.name).name
    repository_file = HERE / filename
elif names:
    repository_file = source_options[selected_name]
    filename = repository_file.name
    raw = repository_file.read_bytes()
else:
    st.info("Upload a five-column Excel workbook to begin.")
    st.stop()

if len(raw) > MAX_UPLOAD:
    st.error("Workbook exceeds the 15 MB safety limit.")
    st.stop()
if not raw:
    st.error("The selected workbook is empty.")
    st.stop()

identity = hashlib.sha256(raw).hexdigest() + ":" + filename
try:
    if st.session_state.get("editor_source") != identity:
        st.session_state["editor_frame"] = _read_rows(raw)
        st.session_state["editor_source"] = identity
    st.caption("You can paste several cells from Excel. Add or remove rows as needed.")
    edited = st.data_editor(
        st.session_state["editor_frame"],
        key=f"crossword_data_{identity[:12]}",
        num_rows="dynamic", hide_index=True, width="stretch",
        column_config={
            "clue": st.column_config.TextColumn(
                "Clue", width="large", help="Question or short description"),
            "answer": st.column_config.TextColumn(
                "Answer", width="medium", help="A–Z letters, spaces or hyphens"),
            "priority": st.column_config.TextColumn(
                "Priority", width="small", help="1 (high), 2, 3 (low)"),
            "lesson": st.column_config.TextColumn("Lesson", width="small"),
            "module": st.column_config.TextColumn("Module", width="small"),
        },
    )
    edited_bytes = _save_rows(raw, edited)
except (ValueError, TypeError, IndexError, KeyError) as error:
    st.error(f"Question bank error: {error}")
    st.stop()

col_save, col_download = st.columns(2)
with col_save:
    if st.button("Save edits to workbook in Codespaces", type="secondary"):
        try:
            # Never allow a browser-provided filename to escape tools/puzzles/.
            if not repository_file.resolve().is_relative_to(HERE.resolve()):
                raise ValueError("Choose a workbook within tools/puzzles.")
            repository_file.write_bytes(edited_bytes)
            st.success(
                f"Saved tools/puzzles/{repository_file.name}. "
                "To preserve the changes on GitHub, commit and push them later."
            )
            # Replace the old editor base on next rerun, preserving the changes.
            st.session_state.pop("editor_source", None)
        except OSError as error:
            st.error(f"Cannot save Excel workbook: {error}")
with col_download:
    st.download_button(
        "Download edited Excel workbook",
        data=edited_bytes,
        file_name=filename, mime=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )

# Use the edited workbook to populate selectors, even before it is saved.
st.subheader("2. Select questions")
temp_for_validation = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
try:
    temp_for_validation.write(edited_bytes)
    temp_for_validation.close()
    from crossword_builder.core import read_workbook
    entries, _ = read_workbook(Path(temp_for_validation.name))
except (ValueError, OSError) as error:
    st.warning(f"Finish or correct your question rows before generating: {error}")
    st.stop()
finally:
    Path(temp_for_validation.name).unlink(missing_ok=True)

scope = st.radio(
    "What should this puzzle review?",
    ["Lesson", "Module", "All"], horizontal=True,
)
mods = sorted({entry.module for entry in entries})
mod = None
les = None
if scope != "All":
    mod = st.selectbox("Module", mods, format_func=lambda x: f"Module {x}")
    if scope == "Lesson":
        lessons = sorted({e.lesson for e in entries if e.module == mod})
        les = st.selectbox("Lesson", lessons, format_func=lambda x: f"Lesson {x}")
selected_count = sum(
    True for e in entries
    if scope == "All" or (e.module == mod and
                          (scope == "Module" or e.lesson == les))
)
st.caption(f"{selected_count} source questions in the selected scope.")

with st.expander("Optional puzzle settings"):
    maximum = st.slider(
        "Maximum answers to fit", min_value=12, max_value=45,
        value=36, step=1,
        help="The program tries this many, then fewer only if needed to fit "
             "the one-page minimum print standards.",
    )
    attempts = st.slider(
        "Layout-search effort", min_value=100, max_value=1000,
        value=300, step=50,
        help="Higher numbers may find better crossword arrangements, but take longer.",
    )
    seed = st.number_input(
        "Reproducible layout seed", min_value=0, max_value=1_000_000,
        value=1, step=1,
    )

if st.button("Generate five crossword choices", type="primary"):
    args = [
        "(workbook)", "--max-words", str(maximum),
        "--attempts", str(attempts), "--seed", str(seed),
        "--top", "5",
    ]
    if scope == "All":
        args.append("--all")
    elif scope == "Module":
        args.extend(["--module", str(mod)])
    else:
        args.extend(["--module", str(mod), "--lesson", str(les)])
    prior = st.session_state.pop("crossword_results", None)
    if prior:
        shutil.rmtree(prior["parent"], ignore_errors=True)
    parent = Path(tempfile.mkdtemp(prefix="fnd-crossword-workshop-"))
    workbook = parent / "questions.xlsx"
    workbook.write_bytes(edited_bytes)
    output = parent / "puzzles"
    args[0] = str(workbook)
    args.extend(["--out", str(output)])
    stdout, stderr = StringIO(), StringIO()
    with st.spinner("Building and checking printable crossword layouts…"):
        try:
            with redirect_stdout(stdout), redirect_stderr(stderr):
                exit_code = generate_cli(args)
            if exit_code:
                st.error(
                    "The crossword could not be generated. " +
                    (stderr.getvalue() or stdout.getvalue())
                )
                shutil.rmtree(parent, ignore_errors=True)
            else:
                report = json.loads((output / "report.json").read_text(
                    encoding="utf-8"
                ))
                st.session_state["crossword_results"] = {
                    "report": report, "path": str(output), "parent": str(parent),
                    "workbook_hash": hashlib.sha256(edited_bytes).hexdigest(),
                }
                st.success(f"{len(report['candidates'])} layouts are ready to compare.")
        except (OSError, ValueError) as error:
            shutil.rmtree(parent, ignore_errors=True)
            st.error(f"Generation failed: {error}")

# 3. Show the FULL printed page, not just a small grid, to judge white space.
saved = st.session_state.get("crossword_results")
if saved:
    st.subheader("3. Compare and choose")
    current_hash = hashlib.sha256(edited_bytes).hexdigest()
    if current_hash != saved["workbook_hash"]:
        st.info("These previews were generated before your latest edits. "
                "Click Generate again when you want to include them.")
    report = saved["report"]
    output = Path(saved["path"])
    options = report["candidates"]
    if len(options) < 5:
        st.info("Fewer than five distinct puzzles met the one-page printing rules.")
    columns = st.columns(2)
    for index, candidate in enumerate(options):
        with columns[index % 2]:
            with st.container(border=True):
                print_info = candidate["print"]
                st.markdown(
                    f"**Option {candidate['rank']} — score "
                    f"{candidate['score']:.1f}/100**"
                )
                st.caption(
                    f"{candidate['metrics']['words']} answers · "
                    f"{candidate['metrics']['words_with_2plus_crossings']} "
                    f"with 2+ crossings · "
                    f"{print_info['clue_font_pt']:g}pt clues · "
                    f"{print_info['square_pt']:g}pt cells"
                )
                st.image(_thumbnail(output / candidate["pdfs"]["normal"]))
                st.caption(
                    f"{candidate['metrics']['borrowed_words']} borrowed answers; "
                    f"{len(candidate['unused_selected_questions'])} unused "
                    "questions from the selected scope."
                )

    selected = st.radio(
        "Which puzzle should be your final version?",
        [c["rank"] for c in options],
        format_func=lambda number: (
            f"Option {number} — {options[number - 1]['metrics']['words']} words"
        ),
        horizontal=True,
    )
    choice = options[selected - 1]
    st.subheader("4. Give your chosen crossword a title")
    st.caption(
        "This is the last step. The title stays on one line at the actual "
        "15-point bold PDF font size. We'll tell you if it needs shortening."
    )
    chosen_title = st.text_input(
        "Final printed title",
        value=report["title"],
        key=f"crossword_title_{saved['workbook_hash'][:10]}_{selected}",
        max_chars=120,
        help="The maximum depends on the measured width of your title, "
             "not a fixed number of characters. Wider letters take more space.",
    )
    try:
        layout = restore_layout(choice)
        measured, permitted = title_measure(chosen_title, layout)
        st.caption(
            f"Title line: {measured:.0f} of {permitted:.0f} points "
            "at 15 pt bold."
        )
        title = validate_title(chosen_title, layout)
    except (ValueError, KeyError, TypeError) as error:
        st.warning(f"Shorten or adjust the title before downloading: {error}")
    else:
        zip_name, named_files, zip_bytes = _finished_pdf_bundle(
            json.dumps(choice, sort_keys=True),
            json.dumps(report["selection"], sort_keys=True),
            title,
        )
        st.download_button(
            "Download chosen puzzle — Normal, Easy and Answer Key (ZIP)",
            data=zip_bytes, file_name=zip_name,
            mime="application/zip", type="primary",
        )
        formats = [("Normal", "normal"), ("Easy", "easy"),
                   ("Answer key", "key")]
        for col, (name, variant) in zip(st.columns(3), formats):
            with col:
                pdf_filename, pdf_data = named_files[variant]
                st.download_button(
                    f"Download {name} PDF",
                    data=pdf_data, file_name=pdf_filename,
                    mime="application/pdf",
                )
    with st.expander("View questions included, borrowed, and omitted"):
        st.dataframe(pd.DataFrame(choice["included"]), hide_index=True)
        if choice["unused_selected_questions"]:
            st.write("Selected questions not used in this layout:")
            st.dataframe(pd.DataFrame(choice["unused_selected_questions"]),
                         hide_index=True)

st.divider()
st.caption(
    "Your workbook and generated files stay within this Codespace until "
    "you save/commit or download them. Keep the forwarded Codespaces port "
    "private. The website's public course pages are not changed."
)
