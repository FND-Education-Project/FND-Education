"""Word selection, legal placement, and layout ranking for educational crosswords.

This module has no network access and does not modify source workbooks. Each
placement must meet the crossword's non-overlap, adjacency, and connectivity
rules. Random seeds make searches reproducible.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from collections import Counter, defaultdict
from pathlib import Path
import random
import re
from typing import Iterable

from openpyxl import load_workbook


HEADERS = ("clue", "answer", "priority", "lesson", "module")
DIRECTIONS = ("across", "down")


@dataclass(frozen=True)
class Entry:
    """One immutable source question, retaining its source worksheet row."""
    clue: str
    original_answer: str
    answer: str
    priority: int
    lesson: int
    module: int
    source_row: int


@dataclass(frozen=True)
class Placed:
    """An answer positioned on the unbounded working grid."""
    entry: Entry
    row: int
    col: int
    direction: str


@dataclass
class Layout:
    """A connected crossword with a measurable learning-oriented quality."""
    placed: list[Placed] = field(default_factory=list)
    cells: dict[tuple[int, int], str] = field(default_factory=dict)
    directions: dict[tuple[int, int], set[str]] = field(default_factory=dict)
    score: float = 0
    metrics: dict[str, float | int] = field(default_factory=dict)

    def bounds(self) -> tuple[int, int, int, int]:
        if not self.cells:
            return (0, 0, -1, -1)
        rows, cols = zip(*self.cells.keys())
        return min(rows), min(cols), max(rows), max(cols)

    def size(self) -> tuple[int, int]:
        a, b, c, d = self.bounds()
        return (c - a + 1, d - b + 1)

    def normalized(self) -> "Layout":
        """Move a puzzle to (0,0) without changing its relationships."""
        row0, col0, _, _ = self.bounds()
        return Layout(
            placed=[Placed(p.entry, p.row - row0, p.col - col0, p.direction)
                    for p in self.placed],
            cells={(r - row0, c - col0): ch for (r, c), ch in self.cells.items()},
            directions={(r - row0, c - col0): set(ds)
                        for (r, c), ds in self.directions.items()},
            score=self.score,
            metrics=dict(self.metrics),
        )

    def signature(self) -> tuple:
        """Ignore translations, but retain answer identities and geometry."""
        n = self.normalized()
        return tuple(sorted((p.entry.answer, p.row, p.col, p.direction)
                            for p in n.placed))

    def numbered(self) -> list[tuple[int, Placed]]:
        """Standard shared numbering for Across/Down at identical starts."""
        numbers: dict[tuple[int, int], int] = {}
        starts = sorted({(p.row, p.col) for p in self.placed})
        for i, start in enumerate(starts, 1):
            numbers[start] = i
        return sorted(
            ((numbers[(p.row, p.col)], p) for p in self.placed),
            key=lambda item: (item[1].direction != "across", item[0]),
        )


def _index(value: object, column: str, row: int) -> int:
    """Parse numbers, 'Module 2', or 'Lesson 03' without lexical misordering."""
    if isinstance(value, bool):
        raise ValueError(f"Row {row}: invalid {column}: {value!r}")
    if isinstance(value, (float, int)):
        if int(value) == value and value > 0:
            return int(value)
    elif isinstance(value, str):
        token = value.strip()
        match = re.fullmatch(r"(?:module|lesson)?[\s_#:-]*(\d+)", token, re.I)
        if match and int(match.group(1)) > 0:
            return int(match.group(1))
    raise ValueError(f"Row {row}: {column} must be a positive number (e.g., 2).")


def _priority(value: object, row: int) -> int:
    """1 is most important; empty entries default to middle priority 2."""
    if value is None or str(value).strip() == "":
        return 2
    names = {"high": 1, "medium": 2, "normal": 2, "low": 3}
    if isinstance(value, str) and value.strip().lower() in names:
        return names[value.strip().lower()]
    try:
        number = int(str(value).strip())
        if number in (1, 2, 3) and float(value) == number:
            return number
    except (TypeError, ValueError):
        pass
    raise ValueError(f"Row {row}: priority must be 1/2/3 or High/Medium/Low.")


def read_workbook(path: Path) -> tuple[list[Entry], list[str]]:
    """Read the first sheet of a five-column .xlsx book; validate every row."""
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Input must be an .xlsx Excel workbook.")
    if not path.is_file() or path.stat().st_size > 15_000_000:
        raise ValueError("Workbook is missing or exceeds the 15 MB limit.")

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.worksheets[0]
        if sheet.max_row and sheet.max_row > 10001:
            raise ValueError("Workbook exceeds 10,000 question rows.")
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("Workbook has no header row.")
        fields = {str(cell).strip().lower(): col for col, cell in enumerate(header)
                  if cell is not None}
        missing = set(HEADERS) - set(fields)
        if missing:
            raise ValueError("Missing column(s): " + ", ".join(sorted(missing)))
        warnings: list[str] = []
        entries: list[Entry] = []
        for rownum, raw in enumerate(rows, 2):
            if not any(value is not None and str(value).strip() for value in raw):
                continue
            def get(name: str) -> object:
                col = fields[name]
                return raw[col] if col < len(raw) else None

            clue = get("clue")
            answer = get("answer")
            if not isinstance(clue, str) or not clue.strip():
                raise ValueError(f"Row {rownum}: clue must contain text.")
            if answer is None or not str(answer).strip():
                raise ValueError(f"Row {rownum}: answer is empty.")
            original = str(answer).strip()
            # Space and hyphen separation are conventional in clue answers.
            # Anything else is rejected rather than silently converted.
            canonical = re.sub(r"[\s\-]+", "", original).upper()
            if not re.fullmatch(r"[A-Z]{2,}", canonical):
                raise ValueError(
                    f"Row {rownum}: answer {original!r} must contain at least "
                    "two A-Z letters (spaces/hyphens permitted)."
                )
            if canonical != original:
                warnings.append(
                    f"Row {rownum}: grid answer {canonical} represents {original!r}."
                )
            entries.append(Entry(
                clue.strip(), original, canonical, _priority(get("priority"), rownum),
                _index(get("lesson"), "lesson", rownum),
                _index(get("module"), "module", rownum), rownum,
            ))
        if not entries:
            raise ValueError("Workbook contains no question rows.")
        return entries, warnings
    finally:
        book.close()


def select_entries(
    entries: Iterable[Entry], mode: str, module: int | None = None,
    lesson: int | None = None,
) -> tuple[list[Entry], list[Entry], list[str]]:
    """Select main pool and, for lessons only, eligible earlier-lesson loans."""
    items = list(entries)
    if mode == "lesson":
        if module is None or lesson is None:
            raise ValueError("Lesson selection requires both --module and --lesson.")
        primary = [e for e in items if (e.module, e.lesson) == (module, lesson)]
        loans = [e for e in items if (e.module, e.lesson) < (module, lesson)]
        loans.sort(key=lambda e: (-e.module, -e.lesson, e.priority, e.source_row))
    elif mode == "module":
        if module is None:
            raise ValueError("Module selection requires --module.")
        primary = [e for e in items if e.module == module]
        loans = []
    elif mode == "all":
        primary, loans = items, []
    else:
        raise ValueError(f"Unknown mode {mode!r}.")
    if not primary:
        raise ValueError("No questions match that selection.")

    # A crossword cannot assign two independent clues to the same answer.
    # Keep the highest-priority primary occurrence; report excluded duplicates.
    selected: dict[str, Entry] = {}
    notes: list[str] = []
    for item in sorted(primary, key=lambda e: (e.priority, e.source_row)):
        if item.answer in selected:
            notes.append(
                f"Duplicate answer {item.answer!r}: row {item.source_row} excluded "
                f"in favour of row {selected[item.answer].source_row}."
            )
        else:
            selected[item.answer] = item
    loans_unique: list[Entry] = []
    seen = set(selected)
    for entry in loans:
        if entry.answer not in seen:
            loans_unique.append(entry)
            seen.add(entry.answer)
    return list(selected.values()), loans_unique, notes


def _points(p: Placed) -> list[tuple[int, int]]:
    if p.direction == "across":
        return [(p.row, p.col + i) for i in range(len(p.entry.answer))]
    return [(p.row + i, p.col) for i in range(len(p.entry.answer))]


def _legal(layout: Layout, entry: Entry, row: int, col: int,
           direction: str, max_side: int,
           existing_bounds: tuple[int, int, int, int] | None = None
           ) -> tuple[int, list[int]] | None:
    """Check intersections, boundaries, and accidental side-by-side letters."""
    dr, dc = (0, 1) if direction == "across" else (1, 0)
    word = entry.answer
    before = (row - dr, col - dc)
    after = (row + len(word) * dr, col + len(word) * dc)
    if before in layout.cells or after in layout.cells:
        return None
    intersection_indices: list[int] = []
    for i, char in enumerate(word):
        pos = (row + dr * i, col + dc * i)
        existing = layout.cells.get(pos)
        if existing is not None:
            if existing != char or direction in layout.directions[pos]:
                return None
            intersection_indices.append(i)
        else:
            near = ((pos[0] - 1, pos[1]), (pos[0] + 1, pos[1])) \
                if direction == "across" else \
                ((pos[0], pos[1] - 1), (pos[0], pos[1] + 1))
            if any(adjacent in layout.cells for adjacent in near):
                return None
    if layout.cells and not intersection_indices:
        return None
    # Avoid re-copying every occupied cell for thousands of candidate
    # placements. A word's end points plus the existing bounding rectangle
    # fully determine whether it fits the maximum grid dimensions.
    end_row = row + dr * (len(word) - 1)
    end_col = col + dc * (len(word) - 1)
    if layout.cells:
        r0, c0, r1, c1 = existing_bounds or layout.bounds()
        new_height = max(r1, row, end_row) - min(r0, row, end_row) + 1
        new_width = max(c1, col, end_col) - min(c0, col, end_col) + 1
    else:
        new_height = abs(end_row - row) + 1
        new_width = abs(end_col - col) + 1
    if new_height > max_side or new_width > max_side:
        return None
    return len(intersection_indices), intersection_indices


def _place(layout: Layout, entry: Entry, row: int, col: int,
           direction: str) -> None:
    placed = Placed(entry, row, col, direction)
    layout.placed.append(placed)
    for pos, char in zip(_points(placed), entry.answer):
        layout.cells[pos] = char
        layout.directions.setdefault(pos, set()).add(direction)


def _best_placement(layout: Layout, entry: Entry, rng: random.Random,
                    max_side: int) -> tuple[int, int, str] | None:
    """Search legal crossing positions, favouring intersections and compactness."""
    by_letter: dict[str, list[tuple[int, int]]] = defaultdict(list)
    for coord, letter in layout.cells.items():
        by_letter[letter].append(coord)
    options: list[tuple[float, int, int, str]] = []
    visited = set()
    y0, x0, y1, x1 = layout.bounds()
    old_r, old_c = y1 - y0 + 1, x1 - x0 + 1
    existing_bounds = (y0, x0, y1, x1)
    for i, char in enumerate(entry.answer):
        for y, x in by_letter.get(char, []):
            for direction in DIRECTIONS:
                row = y if direction == "across" else y - i
                col = x - i if direction == "across" else x
                candidate = (row, col, direction)
                if candidate in visited:
                    continue
                visited.add(candidate)
                result = _legal(
                    layout, entry, row, col, direction, max_side,
                    existing_bounds=existing_bounds
                )
                if result is None:
                    continue
                n, indices = result
                dr, dc = (0, 1) if direction == "across" else (1, 0)
                new_rows = max(y1, row + dr * (len(entry.answer) - 1)) - min(y0, row) + 1
                new_cols = max(x1, col + dc * (len(entry.answer) - 1)) - min(x0, col) + 1
                span = ((max(indices) - min(indices)) / max(1, len(entry.answer)-1)) \
                    if len(indices) > 1 else 0.0
                desirability = (
                    12 * n + 3 * span - .06 * (new_rows * new_cols - old_r * old_c)
                    + rng.random() * 5
                )
                options.append((desirability, row, col, direction))
    if not options:
        return None
    options.sort(reverse=True)
    _, y, x, direction = rng.choice(options[:min(3, len(options))])
    return y, x, direction


def build_attempt(primary: list[Entry], loans: list[Entry], rng: random.Random,
                  max_words: int, max_side: int) -> Layout:
    """Use priority-biased starts, several insertion passes, then earlier lessons."""
    result = Layout()
    pool = list(primary)
    rng.shuffle(pool)
    # Priority and long words help early placement, but jitter avoids one
    # deterministic greedy outcome on every attempt.
    pool.sort(key=lambda e: (e.priority * 4 - min(len(e.answer), 14) / 6
                              + rng.uniform(-3, 3)))
    for entry in pool:
        if len(entry.answer) <= max_side:
            _place(result, entry, 0, 0, rng.choice(DIRECTIONS))
            break
    if not result.placed:
        return result
    used = {result.placed[0].entry.answer}
    for _pass in range(3):
        unplaced = [e for e in pool if e.answer not in used]
        rng.shuffle(unplaced)
        for entry in unplaced:
            if len(result.placed) >= max_words:
                break
            found = _best_placement(result, entry, rng, max_side)
            if found:
                _place(result, entry, *found)
                used.add(entry.answer)
        if len(result.placed) >= max_words:
            break

    # Borrow from an earlier lesson only to complete a lesson puzzle after
    # the target lesson's answer pool has been attempted.
    if len(result.placed) < max_words and loans:
        older = sorted(loans, key=lambda e: (
            -e.module, -e.lesson, e.priority, rng.random()))
        for entry in older:
            if len(result.placed) >= max_words:
                break
            if len(entry.answer) > max_side:
                continue
            found = _best_placement(result, entry, rng, max_side)
            if found:
                _place(result, entry, *found)
    return result.normalized()


def evaluate(layout: Layout, primary: list[Entry], max_words: int) -> None:
    """Rank educational value without weakening the mandatory placement rules."""
    if not layout.placed:
        layout.score = 0
        return
    main = {e.answer for e in primary}
    lengths = [len(p.entry.answer) for p in layout.placed]
    crossing = []
    spreads = []
    for p in layout.placed:
        opposite = "down" if p.direction == "across" else "across"
        indices = [i for i, pos in enumerate(_points(p))
                   if opposite in layout.directions[pos]]
        crossing.append(len(indices))
        if len(indices) > 1:
            spreads.append((max(indices) - min(indices)) / max(1, len(p.entry.answer)-1))
        else:
            spreads.append(0.0)
    count = len(layout.placed)
    height, width = layout.size()
    filled_ratio = len(layout.cells) / max(1, width * height)
    across = sum(p.direction == "across" for p in layout.placed)
    # Reward 2+ crossings for mid-length words and 3+ for long answers.
    crossing_score = sum(
        min(1.0, n / (3 if length >= 10 else 2 if length >= 6 else 1))
        for n, length in zip(crossing, lengths)
    ) / count
    coverage = min(1.0, count / max(1, min(max_words, len(primary))))
    main_ratio = sum(p.entry.answer in main for p in layout.placed) / count
    importance = sum((4 - p.entry.priority) / 3 for p in layout.placed) / count
    balance = 1 - abs(2 * across - count) / count
    layout.score = round(
        30 * coverage + 29 * crossing_score + 12 * sum(spreads) / count
        + 9 * filled_ratio + 5 * balance + 7 * importance + 8 * main_ratio,
        2,
    )
    layout.metrics = {
        "words": count,
        "main_lesson_words": sum(p.entry.answer in main for p in layout.placed),
        "borrowed_words": sum(p.entry.answer not in main for p in layout.placed),
        "words_with_2plus_crossings": sum(n >= 2 for n in crossing),
        "long_words_with_3plus_crossings": sum(
            n >= 3 and length >= 10 for n, length in zip(crossing, lengths)),
        "intersection_letters": sum(crossing) // 2,
        "rows": height, "columns": width,
        "across": across, "down": count - across,
        "coverage_percent": round(100 * coverage, 1),
    }


def top_layouts(primary: list[Entry], loans: list[Entry], *,
                attempts: int = 250, top: int = 5,
                max_words: int = 24, max_side: int = 29,
                seed: int = 1, progress=None) -> list[Layout]:
    """Generate distinct, reproducible, quality-ranked connected layouts."""
    if not 3 <= max_words <= 60 or not 8 <= max_side <= 40:
        raise ValueError("Require 3 <= max_words <= 60 and 8 <= max_side <= 40.")
    if not 1 <= attempts <= 5000 or not 1 <= top <= 20:
        raise ValueError("Require 1 <= attempts <= 5000 and 1 <= top <= 20.")
    rng = random.Random(seed)
    best: dict[tuple, Layout] = {}
    for attempt_index in range(attempts):
        layout = build_attempt(primary, loans, rng, max_words, max_side)
        if len(layout.placed) >= 3:
            evaluate(layout, primary, max_words)
            signature = layout.signature()
            if signature not in best or layout.score > best[signature].score:
                best[signature] = layout
        # Avoid flooding browser updates on every candidate.
        completed = attempt_index + 1
        if progress is not None and (
            completed == attempts or completed == 1 or completed % 4 == 0
        ):
            progress(completed, attempts)
    return sorted(best.values(), key=lambda x: (
        -x.score, -int(x.metrics["main_lesson_words"]),
        -int(x.metrics["words"]), x.signature(),
    ))[:max(top * 8, top)]


def verify_layout(layout: Layout) -> None:
    """Independent reconstruction detects clashes, adjacency and bad metadata."""
    check = Layout()
    if len({p.entry.answer for p in layout.placed}) != len(layout.placed):
        raise ValueError("Duplicate answer in layout.")
    for p in layout.placed:
        legal = _legal(check, p.entry, p.row, p.col, p.direction, 1000)
        if legal is None:
            raise ValueError(f"Invalid crossword placement for {p.entry.answer}.")
        _place(check, p.entry, p.row, p.col, p.direction)
    if check.cells != layout.cells or check.directions != layout.directions:
        raise ValueError("Stored letters/orientations do not match placements.")
